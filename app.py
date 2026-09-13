# 영수증 OCR 자동 정리 웹앱 (Streamlit)
# - 영수증 이미지를 업로드하면 Upstage OCR로 텍스트를 추출하고,
#   Upstage LLM(JSON Schema)으로 가게이름/구매일시/물품/총구매금액을 구조화하여 표로 보여줌
# - 결과는 엑셀(.xlsx)로 다운로드 가능 (파일명·가게이름 셀 병합 포함)

import io
import json
import os

import openpyxl
import pandas as pd
import requests
import streamlit as st
from dotenv import load_dotenv
from openai import OpenAI
from openpyxl.styles import Alignment

load_dotenv()  # 로컬 개발 시 .env 파일을 읽음 (배포 환경에서는 st.secrets 사용)


def get_config(key, default=None):
    """배포 환경(st.secrets)과 로컬 환경(.env) 양쪽을 모두 지원"""
    try:
        if key in st.secrets:
            return st.secrets[key]
    except Exception:
        pass
    return os.environ.get(key, default)


UPSTAGE_API_KEY = get_config("UPSTAGE_API_KEY")
UPSTAGE_BASE_URL = get_config("UPSTAGE_BASE_URL", "https://api.upstage.ai/v1")

OCR_URL = "https://api.upstage.ai/v1/document-digitization"

RECEIPT_SCHEMA = {
    "type": "object",
    "properties": {
        "store_name": {"type": "string", "description": "가게(매장) 이름"},
        "purchase_datetime": {"type": "string", "description": "구매 일시"},
        "items": {
            "type": "array",
            "description": "구매한 물품 목록",
            "items": {
                "type": "object",
                "properties": {
                    "item_name": {"type": "string", "description": "수량, 단가 표기를 제외한 순수 물품명"},
                    "quantity": {"type": "number", "description": "구매 수량"},
                    "unit_price": {"type": "number", "description": "물품 1개당 단가"},
                },
                "required": ["item_name", "quantity", "unit_price"],
                "additionalProperties": False,
            },
        },
        "total_amount": {"type": "number", "description": "총 구매금액"},
    },
    "required": ["store_name", "purchase_datetime", "items", "total_amount"],
    "additionalProperties": False,
}

RESULT_COLUMNS = ["파일명", "가게이름", "구매일시", "물품명", "수량", "단가", "금액", "총 구매금액"]


def extract_receipt_rows(client, file_bytes, filename):
    """영수증 이미지(메모리 바이트) 1개를 OCR + LLM으로 처리하여 물품별 행 리스트를 반환"""
    headers = {"Authorization": f"Bearer {UPSTAGE_API_KEY}"}
    ocr_response = requests.post(
        OCR_URL,
        headers=headers,
        files={"document": (filename, io.BytesIO(file_bytes))},
        data={"model": "ocr"},
        timeout=60,
    )
    ocr_response.raise_for_status()
    ocr_text = ocr_response.json()["text"]

    llm_response = client.chat.completions.create(
        model="solar-mini",
        timeout=60,
        messages=[
            {
                "role": "system",
                "content": "너는 영수증 OCR 텍스트에서 정보를 정확하게 추출하는 도우미야. "
                "물품명(item_name)에는 수량이나 가격 표기를 포함하지 말고 순수 상품명만 넣어.",
            },
            {"role": "user", "content": ocr_text},
        ],
        response_format={
            "type": "json_schema",
            "json_schema": {"name": "receipt_info", "schema": RECEIPT_SCHEMA, "strict": True},
        },
    )
    info = json.loads(llm_response.choices[0].message.content)

    return [
        {
            "파일명": filename,
            "가게이름": info["store_name"],
            "구매일시": info["purchase_datetime"],
            "물품명": item["item_name"],
            "수량": item["quantity"],
            "단가": item["unit_price"],
            "금액": round(item["quantity"] * item["unit_price"], 2),
            "총 구매금액": info["total_amount"],
        }
        for item in info["items"]
    ]


def merge_by_group(ws, group_col_idx, target_col_idxs, n_rows, header_row=1):
    """group_col_idx 값이 연속으로 같은 구간을 찾아 target_col_idxs 열들을 함께 병합"""
    start_row = header_row + 1
    last_row = header_row + n_rows
    merge_start = start_row
    prev_key = ws.cell(row=start_row, column=group_col_idx).value
    for row in range(start_row + 1, last_row + 2):
        current_key = ws.cell(row=row, column=group_col_idx).value if row <= last_row else object()
        if current_key != prev_key:
            if row - 1 > merge_start:
                for col_idx in target_col_idxs:
                    ws.merge_cells(start_row=merge_start, start_column=col_idx, end_row=row - 1, end_column=col_idx)
                    ws.cell(row=merge_start, column=col_idx).alignment = Alignment(vertical="center")
            merge_start = row
            prev_key = current_key


def build_merged_excel_bytes(df):
    """DataFrame을 파일명/가게이름 셀 병합이 적용된 엑셀 바이트로 변환 (디스크에 저장하지 않음)"""
    buffer = io.BytesIO()
    df.to_excel(buffer, index=False)
    buffer.seek(0)

    wb = openpyxl.load_workbook(buffer)
    ws = wb.active
    merge_by_group(ws, group_col_idx=1, target_col_idxs=[1, 2], n_rows=len(df))

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out.getvalue()


# ==============================
# Streamlit UI
# ==============================
st.set_page_config(page_title="영수증 OCR 자동 정리", page_icon="🧾")
st.title("🧾 영수증 OCR 자동 정리")
st.caption("영수증 이미지를 업로드하면 Upstage OCR + LLM으로 항목을 자동 추출해 표로 정리합니다.")

if not UPSTAGE_API_KEY:
    st.error("UPSTAGE_API_KEY가 설정되어 있지 않습니다. .env 파일 또는 배포 환경의 Secrets를 확인하세요.")
    st.stop()

uploaded_files = st.file_uploader(
    "영수증 이미지를 업로드하세요 (여러 장 가능)",
    type=["jpg", "jpeg", "png", "webp", "bmp"],
    accept_multiple_files=True,
)

if st.button("분석하기", type="primary", disabled=not uploaded_files):
    client = OpenAI(api_key=UPSTAGE_API_KEY, base_url=UPSTAGE_BASE_URL)

    all_rows = []
    progress = st.progress(0.0)
    status_area = st.empty()
    for i, uploaded_file in enumerate(uploaded_files):
        status_area.info(f"({i + 1}/{len(uploaded_files)}) {uploaded_file.name} 처리 중...")
        try:
            rows = extract_receipt_rows(client, uploaded_file.getvalue(), uploaded_file.name)
            all_rows.extend(rows)
        except Exception as e:
            st.warning(f"{uploaded_file.name} 처리 실패: {e}")
        progress.progress((i + 1) / len(uploaded_files))
    status_area.empty()

    if all_rows:
        st.session_state["result_df"] = pd.DataFrame(all_rows, columns=RESULT_COLUMNS)
        st.success(f"영수증 {len(uploaded_files)}개 처리 완료!")

if "result_df" in st.session_state:
    result_df = st.session_state["result_df"]
    st.dataframe(result_df, use_container_width=True)

    excel_bytes = build_merged_excel_bytes(result_df)
    st.download_button(
        label="엑셀로 다운로드",
        data=excel_bytes,
        file_name="receipt_result.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
